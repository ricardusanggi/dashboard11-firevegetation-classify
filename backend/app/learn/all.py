# ===========================
# BLOK 1 / 3
# Dataset, Parameter Manager, Utilities, Excel Writer
# ===========================

import datetime
import os
import random
import time
import json
from itertools import product
from pathlib import Path
from typing import List, Dict, Tuple, Any

import numpy as np
import cv2
from scipy import ndimage
from PIL import Image
import matplotlib.pyplot as plt

import torch
from torch.utils.data import Dataset, DataLoader, Subset
import torchvision.transforms as T

from openpyxl import Workbook


# -----------------------------
# Device helper
# -----------------------------
def get_device() -> torch.device:
    if torch.cuda.is_available():
        print("🔥 Using CUDA GPU")
        return torch.device("cuda")
    if hasattr(torch.backends, "mps") and torch.backends.mps.is_available():
        print("🍏 Using Apple MPS")
        return torch.device("mps")
    print("⚙️ Using CPU")
    return torch.device("cpu")


# ==========================
#       DATASET
# ==========================
class SegDataset(Dataset):
    VALID_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}

    def __init__(self, img_dir, mask_dir, img_size=(256,256), augment=False):
        self.img_dir = str(img_dir)
        self.mask_dir = str(mask_dir)
        self.img_size = tuple(img_size)
        self.augment = bool(augment)

        if not os.path.isdir(self.img_dir):
            raise FileNotFoundError(self.img_dir)
        if not os.path.isdir(self.mask_dir):
            raise FileNotFoundError(self.mask_dir)

        self.images = sorted([
            f for f in os.listdir(self.img_dir)
            if os.path.splitext(f)[1].lower() in self.VALID_EXT
        ])

        from PIL import Image

        def pil_rgb_to_tensor(pil_img, size):
            pil_img = pil_img.resize(size, Image.BILINEAR)
            data = pil_img.tobytes()
            w, h = pil_img.size
            arr = torch.tensor(list(data), dtype=torch.uint8).reshape(h, w, 3)
            return arr.permute(2,0,1).float() / 255.0

        def pil_mask_to_tensor(pil_img, size):
            pil_img = pil_img.resize(size, Image.NEAREST)
            data = pil_img.tobytes()
            w, h = pil_img.size
            arr = torch.tensor(list(data), dtype=torch.uint8).reshape(h, w)
            return arr.unsqueeze(0).float() / 255.0

        self.t_img  = lambda img: pil_rgb_to_tensor(img, self.img_size)
        self.t_mask = lambda m: pil_mask_to_tensor(m, self.img_size)

    def find_mask_path(self, img_name):
        base = os.path.splitext(img_name)[0]

        for ext in self.VALID_EXT:
            p = os.path.join(self.mask_dir, base + ext)
            if os.path.exists(p):
                return p

        for ext in self.VALID_EXT:
            p = os.path.join(self.mask_dir, base + "_mask" + ext)
            if os.path.exists(p):
                return p

        raise FileNotFoundError(f"Mask not found for {img_name}")

    def _augment_brown(self, img_tensor):
        img = (img_tensor.permute(1,2,0).numpy()*255).astype(np.uint8)
        hsv = cv2.cvtColor(img, cv2.COLOR_RGB2HSV)
        h,s,v = cv2.split(hsv)

        shift = random.randint(5,15)
        h = ((h + shift) % 180).astype(np.uint8)

        hsv2 = cv2.merge([h,s,v])
        img2 = cv2.cvtColor(hsv2, cv2.COLOR_HSV2RGB)

        if random.random() < 0.7:
            muddy = np.array([115,80,45], np.uint8)
            alpha = random.uniform(0.15, 0.35)
            img2 = (img2*(1-alpha) + muddy*alpha).astype(np.uint8)

        if random.random() < 0.5:
            hsv3 = cv2.cvtColor(img2, cv2.COLOR_RGB2HSV)
            hh,ss,vv = cv2.split(hsv3)
            ss = np.clip(ss * random.uniform(0.6, 0.95), 0, 255).astype(np.uint8)
            hsv3 = cv2.merge([hh,ss,vv])
            img2 = cv2.cvtColor(hsv3, cv2.COLOR_HSV2RGB)

        return torch.tensor(img2).permute(2,0,1).float()/255.

    def _augment_mask_morph(self, mask_tensor):
        m = (mask_tensor.numpy()[0]*255).astype(np.uint8)
        kernel = np.ones((5,5), np.uint8)
        op = random.choice(["none","erode","dilate","open","close"])
        if op=="erode":
            m = cv2.erode(m, kernel)
        elif op=="dilate":
            m = cv2.dilate(m, kernel)
        elif op=="open":
            m = cv2.morphologyEx(m, cv2.MORPH_OPEN, kernel)
        elif op=="close":
            m = cv2.morphologyEx(m, cv2.MORPH_CLOSE, kernel)

        return torch.tensor(m/255.).unsqueeze(0).float()

    def __len__(self):
        return len(self.images)

    def __getitem__(self, idx):
        name = self.images[idx]
        img = Image.open(os.path.join(self.img_dir,name)).convert("RGB")
        mask = Image.open(self.find_mask_path(name)).convert("L")

        img_t = self.t_img(img)
        mask_t = (self.t_mask(mask) > 0.5).float()

        if self.augment:
            if random.random() < 0.6:
                img_t = self._augment_brown(img_t)
            if random.random() < 0.4:
                mask_t = self._augment_mask_morph(mask_t)

        return img_t.float(), mask_t.float()


# ------------------------------------------------
# Save overlay prediction image
# ------------------------------------------------
def save_pred_overlay(img_tensor, pred_mask, out_path):

    try:
        if isinstance(img_tensor, torch.Tensor):
            img_np = img_tensor.detach().cpu().permute(1,2,0).clamp(0,1).mul(255).byte().numpy()
        else:
            img_np = np.array(img_tensor)

        if pred_mask.max() <= 1:
            pm = (pred_mask > 0.5).astype(np.uint8) * 255
        else:
            pm = pred_mask.astype(np.uint8)

        pm = pm.reshape(img_np.shape[0], img_np.shape[1])

        mask_rgb = np.stack([pm,pm,pm], axis=-1)

        overlay = img_np.copy()
        red = np.zeros_like(img_np)
        red[...,0] = 255
        alpha = 0.4
        overlay[pm>0] = (overlay[pm>0]*(1-alpha) + red[pm>0]*alpha).astype(np.uint8)

        combined = np.concatenate([img_np, mask_rgb, overlay], axis=1)
        Image.fromarray(combined).save(out_path)

    except Exception as e:
        #print("[SAVE ERROR]", e)


# ------------------------------------------------
# Save Excel Journal
# ------------------------------------------------
def save_excel_journal(out_path,
                       analysis_rows,
                       params_by_run,
                       train_rows_by_run,
                       test_rows_by_run):

    wb = Workbook()
    ws = wb.active
    ws.title = "ANALYSIS"

    header = [
        "run_id","timestamp",
        "img_size","split_ratio","num_train_images","num_test_images",
        "optimizer","learning_rate","epochs","batch_size","loss_fn",
        "bottleneck_filters","unet_depth",
        "use_pso","pso_particles","pso_iterations","pso_w","pso_c1","pso_c2",
        "pso_trigger_every_batches",
        "train_mean_loss","train_mean_iou","train_mean_dice",
        "train_mean_pa","train_time_total","train_time_per_epoch",
        "test_mean_loss","test_mean_iou","test_mean_dice","test_mean_pa",
        "model_path","preds_dir","train_samples_dir"
    ]

    ws.append(header)

    for r in analysis_rows:
        ws.append([r.get(k,"") for k in header])

    for run_id, params in params_by_run.items():
        ws2 = wb.create_sheet(f"PARAMETERS_run_{run_id}")
        ws2.append(["parameter","value"])
        for k,v in params.items():
            ws2.append([k, json.dumps(v) if isinstance(v,(dict,list)) else str(v)])

    for run_id, rows in train_rows_by_run.items():
        ws3 = wb.create_sheet(f"TRAIN_run_{run_id}")
        ws3.append(["epoch","loss","iou","dice","pa","time_s"])
        for r in rows:
            ws3.append(r)

    for run_id, rows in test_rows_by_run.items():
        ws4 = wb.create_sheet(f"TEST_run_{run_id}")
        ws4.append(["image_name","iou","dice","pa","pred_path"])
        for r in rows:
            ws4.append(r)

    wb.save(out_path)
    #print("Excel saved:", out_path)


# ------------------------------------------------
# Build parameter combinations
# ------------------------------------------------
def build_param_combinations(
    optimizer_list,
    lr_list,
    epochs_list,
    batch_list,
    pso_particles_list,
    pso_iterations_list,
    bottleneck_list,
    depth_list,
    max_runs=10):

    combos = list(product(
        optimizer_list,
        lr_list,
        epochs_list,
        batch_list,
        pso_particles_list,
        pso_iterations_list,
        bottleneck_list,
        depth_list
    ))

    configs = []
    for (opt,lr,ep,bs,pso_p,pso_it,bf,d) in combos[:max_runs]:
        configs.append({
            "optimizer": opt,
            "learning_rate": float(lr),
            "epochs": int(ep),
            "batch_size": int(bs),
            "pso_particles": int(pso_p),
            "pso_iterations": int(pso_it),
            "bottleneck_filters": int(bf),
            "unet_depth": int(d)
        })

    return configs

# ===========================
# BLOK 2 / 3
# Model, Loss/Metric, PSO, Training, Evaluation
# ===========================

import torch
import torch.nn as nn
import torch.nn.functional as F
from typing import List, Dict, Any


# ======================================================
#                   UNET MODEL
# ======================================================
class DoubleConv(nn.Module):
    def __init__(self, in_ch, out_ch):
        super().__init__()
        self.conv = nn.Sequential(
            nn.Conv2d(in_ch,out_ch,3,padding=1),
            nn.ReLU(inplace=True),
            nn.Conv2d(out_ch,out_ch,3,padding=1),
            nn.ReLU(inplace=True)
        )
    def forward(self, x):
        return self.conv(x)

class UNetSmall(nn.Module):
    def __init__(self, in_ch=3, base_filters=[64,128], bottleneck_filters=256):
        super().__init__()

        f1,f2 = base_filters

        self.enc1 = DoubleConv(in_ch,f1)
        self.p1   = nn.MaxPool2d(2)

        self.enc2 = DoubleConv(f1,f2)
        self.p2   = nn.MaxPool2d(2)

        self.bottleneck = DoubleConv(f2,bottleneck_filters)

        self.u2 = nn.ConvTranspose2d(bottleneck_filters,f2,2,stride=2)
        self.dec2 = DoubleConv(f2+f2,f2)

        self.u1 = nn.ConvTranspose2d(f2,f1,2,stride=2)
        self.dec1 = DoubleConv(f1+f1,f1)

        self.outc = nn.Conv2d(f1,1,1)

    def forward(self,x):
        x1 = self.enc1(x)
        x2 = self.enc2(self.p1(x1))

        xb = self.bottleneck(self.p2(x2))

        d2 = self.u2(xb)
        d2 = self.dec2(torch.cat([d2,x2],dim=1))

        d1 = self.u1(d2)
        d1 = self.dec1(torch.cat([d1,x1],dim=1))

        return self.outc(d1)


# ======================================================
#               LOSSES & METRICS
# ======================================================
class DiceLoss(nn.Module):
    def forward(self, pred_logits, target):
        pred = torch.sigmoid(pred_logits)
        smooth = 1e-6
        inter = (pred*target).sum()
        union = pred.sum()+target.sum()
        dice = (2*inter+smooth)/(union+smooth)
        return 1 - dice

bce_loss = nn.BCEWithLogitsLoss()
dice_loss = DiceLoss()

def combined_loss(logits, target, alpha=0.5):
    return alpha*bce_loss(logits,target) + (1-alpha)*dice_loss(logits,target)

def tv_loss(prob):
    dy = torch.abs(prob[:,:,1:,:] - prob[:,:,:-1,:]).mean()
    dx = torch.abs(prob[:,:,:,1:] - prob[:,:,:,:-1]).mean()
    return dx+dy

def metric_iou(pred,target):
    pred = (pred>0.5).float()
    target = (target>0.5).float()
    inter = (pred*target).sum().item()
    union = pred.sum().item()+target.sum().item()-inter+1e-6
    return inter/union

def metric_dice(pred,target):
    pred = (pred>0.5).float()
    target = (target>0.5).float()
    inter = (pred*target).sum().item()
    denom = pred.sum().item()+target.sum().item()+1e-6
    return (2*inter)/denom

def metric_pa(pred,target):
    pred = (pred>0.5).float()
    return (pred == (target>0.5).float()).float().mean().item()

# ============================
# LION OPTIMIZER (PyTorch)
# ============================
class Lion(torch.optim.Optimizer):
    """
    Lion Optimizer: Evolved Sign Momentum
    Paper: https://arxiv.org/abs/2302.06675
    """

    def __init__(self, params,
                 lr=1e-4,
                 betas=(0.9, 0.99),
                 weight_decay=0.0):

        defaults = dict(
            lr=lr,
            betas=betas,
            weight_decay=weight_decay
        )
        super(Lion, self).__init__(params, defaults)

    @torch.no_grad()
    def step(self, closure=None):

        for group in self.param_groups:
            lr = group["lr"]
            beta1, beta2 = group["betas"]
            wd = group["weight_decay"]

            for p in group["params"]:
                if p.grad is None:
                    continue

                grad = p.grad

                # Weight decay
                if wd != 0:
                    grad = grad.add(p, alpha=wd)

                # Get state
                state = self.state[p]
                if len(state) == 0:
                    state["exp_avg"] = torch.zeros_like(p)

                exp_avg = state["exp_avg"]

                # Update
                update = exp_avg.mul(beta1).add(grad, alpha=1 - beta1).sign()
                p.add_(update, alpha=-lr)

                # Momentum update
                exp_avg.mul_(beta2).add_(grad, alpha=1 - beta2)

        return None


# ======================================================
#                    PSO CPU
# ======================================================
class PSO_CPU:
    def __init__(self, dim, n_particles=4, w=0.5, c1=1.0, c2=1.0, init_pos=None):
        self.dim = dim
        self.n = n_particles
        if init_pos is None:
            self.pos = torch.randn(self.n,self.dim)*0.01
        else:
            self.pos = init_pos.float().unsqueeze(0).repeat(self.n,1)
            self.pos += torch.randn(self.n,self.dim)*0.01

        self.vel = torch.zeros_like(self.pos)
        self.pbest = self.pos.clone()
        self.pbest_score = torch.full((self.n,), float('inf'))
        self.gbest = self.pos[0].clone()
        self.gbest_score = float('inf')

        self.w=w; self.c1=c1; self.c2=c2

    def step(self, fitness):
        for i in range(self.n):
            fi = float(fitness[i])
            if fi < self.pbest_score[i]:
                self.pbest_score[i] = fi
                self.pbest[i] = self.pos[i].clone()
            if fi < self.gbest_score:
                self.gbest_score = fi
                self.gbest = self.pos[i].clone()

        r1 = torch.rand_like(self.pos)
        r2 = torch.rand_like(self.pos)

        self.vel = (self.w*self.vel
                    + self.c1*r1*(self.pbest - self.pos)
                    + self.c2*r2*(self.gbest - self.pos))

        self.pos = self.pos + self.vel

        return self.gbest.clone(), self.gbest_score

def get_bvec_cpu(model):
    parts=[]
    for p in model.bottleneck.parameters():
        parts.append(p.detach().cpu().view(-1))
    if len(parts)==0:
        return torch.tensor([])
    return torch.cat(parts).float()

def set_bvec_from_cpu(model, vec_cpu):
    idx=0
    for p in model.bottleneck.parameters():
        n = p.numel()
        seg = vec_cpu[idx:idx+n].view_as(p).to(p.device)
        p.data.copy_(seg)
        idx+=n


# ======================================================
#           TRAIN 1 EPOCH
# ======================================================
def train_one_epoch(model, loader, optimizer, device,
                    tv_weight=0.0,
                    pso_obj=None,
                    pso_trigger_every_batches=50,
                    pso_iterations=1,
                    cfg=None):

    if cfg is None: cfg = {}

    model.train()
    total_loss=total_iou=total_dice=total_pa=0.0
    steps=0
    global_step=0

    use_interval = cfg.get("use_pso_interval", True)

    for imgs, masks in loader:

        imgs  = imgs.to(device)
        masks = masks.to(device)

        optimizer.zero_grad()

        logits = model(imgs)
        probs  = torch.sigmoid(logits)

        loss = combined_loss(logits,masks)
        if tv_weight>0:
            loss = loss + tv_weight*tv_loss(probs)

        loss.backward()
        optimizer.step()

        total_loss += loss.item()
        total_iou  += metric_iou(probs.detach(), masks)
        total_dice += metric_dice(probs.detach(), masks)
        total_pa   += metric_pa(probs.detach(), masks)
        steps+=1
        global_step+=1

        # ============= PSO =============
        if pso_obj is not None:

            trigger_now = (global_step % pso_trigger_every_batches == 0) if use_interval else True

            if trigger_now:

                fitness = torch.empty(pso_obj.n)
                real_model = model.module if isinstance(model, nn.DataParallel) else model

                for i in range(pso_obj.n):
                    set_bvec_from_cpu(real_model, pso_obj.pos[i])

                    with torch.no_grad():
                        logits_try = model(imgs)
                        f = combined_loss(logits_try, masks).item()
                    fitness[i] = f

                gbest_vec, gbest_score = pso_obj.step(fitness)

                set_bvec_from_cpu(real_model, gbest_vec)

        # cleanup memory per-batch
        del imgs, masks, logits, probs
        torch.cuda.empty_cache()

    if steps==0:
        return 0,0,0,0,0

    return (total_loss/steps,
            total_iou/steps,
            total_dice/steps,
            total_pa/steps,
            steps)


# ======================================================
#              EVALUATE
# ======================================================
def evaluate_on_loader(model: nn.Module,
                       loader: DataLoader,
                       device: torch.device,
                       postproc: bool = True,
                       binary_threshold: float = 0.5,
                       min_comp_area: int = 300,
                       save_preds_dir: str = None):

    model.eval()
    test_rows = []
    ious, dices, pas = [], [], []

    dataset = loader.dataset
    if isinstance(dataset, Subset):
        base_ds = dataset.dataset
        names = [base_ds.images[i] for i in dataset.indices]
    else:
        names = dataset.images

    import numpy as _np

    for idx, (imgs, masks) in enumerate(loader):
        name = names[idx] if idx < len(names) else f"img_{idx}"

        imgs = imgs.to(device)
        masks = masks.to(device)

        # FORWARD
        with torch.no_grad():
            logits = model(imgs)
            p = torch.sigmoid(logits)[0, 0].cpu()

        pred_bin = (p > binary_threshold).to(torch.uint8)

        # ===========================================
        # AUTO-DETEKSI RUN BURUK
        # ===========================================
        h, w = pred_bin.shape
        num_white = int(pred_bin.sum())
        num_pixels = h * w

        all_black_flag = (num_white == 0)
        all_white_flag = (num_white == num_pixels)

        if all_black_flag or all_white_flag:
            print(f"⚠️ BAD RUN DETECTED ({name}):",
                "ALL BLACK" if all_black_flag else "ALL WHITE")

            # Kembalikan sinyal ke run_config tanpa melakukan evaluasi penuh
            return [], {
                "mean_iou": 0.0,
                "mean_dice": 0.0,
                "mean_pa": 0.0,
                "bad_run": True
            }


        # ============= POST PROCESSING =============
        if postproc:
            arr = (pred_bin * 255).byte().numpy()
            kernel = _np.ones((5,5), _np.uint8)

            clean = cv2.morphologyEx(arr, cv2.MORPH_OPEN, kernel)
            clean = cv2.morphologyEx(clean, cv2.MORPH_CLOSE, kernel)

            fill = ndimage.binary_fill_holes(clean>0).astype(_np.uint8)*255

            nlab, labels = cv2.connectedComponents(fill)
            final = _np.zeros_like(fill)
            for lab in range(1, nlab):
                comp = (labels == lab)
                if comp.sum() > min_comp_area:
                    final[comp] = 255

            final_mask = torch.tensor((final > 0).astype("float32"))
        else:
            final_mask = pred_bin.float()

        # METRICS
        mt = masks[0].cpu().float()

        iou = metric_iou(final_mask, mt)
        dice = metric_dice(final_mask, mt)
        pa = metric_pa(final_mask, mt)

        ious.append(iou)
        dices.append(dice)
        pas.append(pa)

        # ============= SAVE FILES =============
        out_path_original = ""
        out_path_mask = ""
        out_path_pred = ""
        out_path_compare = ""

        if save_preds_dir is not None:
            os.makedirs(save_preds_dir, exist_ok=True)

            # --- ORIGINAL IMAGE (RGB) ---
            img_np = (imgs[0].cpu().clamp(0,1) * 255).byte().permute(1,2,0).numpy()
            pil_img = Image.fromarray(img_np, mode="RGB")
            out_path_original = os.path.join(save_preds_dir, f"{name}_original.png")
            pil_img.save(out_path_original)

            # --- ORIGINAL MASK (convert to grayscale RGB) ---
            mask_np = (mt.squeeze().numpy() * 255).astype(np.uint8)
            pil_mask = Image.fromarray(mask_np, mode="L").convert("RGB")
            out_path_mask = os.path.join(save_preds_dir, f"{name}_mask.png")
            pil_mask.save(out_path_mask)

            # --- PREDICTED MASK ---
            pred_np = (final_mask.squeeze().numpy() * 255).astype(np.uint8)
            pil_pred = Image.fromarray(pred_np, mode="L").convert("RGB")
            out_path_pred = os.path.join(save_preds_dir, f"{name}_pred.png")
            pil_pred.save(out_path_pred)

            # --- 3-PANEL COMPARISON IMAGE ---
            W, H = pil_img.size
            combined = Image.new("RGB", (W*3, H))
            combined.paste(pil_img, (0, 0))      # Original
            combined.paste(pil_mask, (W, 0))     # Ground Truth Mask
            combined.paste(pil_pred, (2*W, 0))   # Predicted Mask

            out_path_compare = os.path.join(save_preds_dir, f"{name}_compare.png")
            combined.save(out_path_compare)

            # --- DIFFERENCE MAP (Correct = grayscale, Wrong = red) ---

            gt_bin = (mt.numpy().squeeze() > 0.5).astype(np.uint8)
            pred_bin2 = (final_mask.numpy().squeeze() > 0.5).astype(np.uint8)

            # Binary ground truth mask → grayscale 0/255
            gray = (gt_bin * 255).astype(np.uint8)

            H, W = gray.shape
            diff = np.zeros((H, W, 3), dtype=np.uint8)

            # Set grayscale where correct (pred == gt)
            correct = (gt_bin == pred_bin2)
            diff[correct] = np.stack([gray[correct],
                                    gray[correct],
                                    gray[correct]], axis=-1)

            # Wrong pixels → red
            wrong = (gt_bin != pred_bin2)
            diff[wrong] = [255, 0, 0]

            pil_diff = Image.fromarray(diff, mode="RGB")



            # --- 4-PANEL COMPARISON IMAGE ---
            W, H = pil_img.size
            combined = Image.new("RGB", (W * 4, H))

            combined.paste(pil_img, (0, 0))          # Original
            combined.paste(pil_mask, (W, 0))         # Ground Truth
            combined.paste(pil_pred, (2 * W, 0))     # Prediction
            combined.paste(pil_diff, (3 * W, 0))     # Difference map

            out_path_compare = os.path.join(save_preds_dir, f"{name}_compare2.png")
            combined.save(out_path_compare)


        # add table row
        test_rows.append([
            name,
            float(iou),
            float(dice),
            float(pa),
            out_path_pred
        ])

        # cleanup
        del imgs, masks, logits
        torch.cuda.empty_cache()

    summary = {
        "mean_iou": float(sum(ious)/len(ious)),
        "mean_dice": float(sum(dices)/len(dices)),
        "mean_pa": float(sum(pas)/len(pas)),
    }

    return test_rows, summary


# ===========================
# BLOK 3 / 3 — DENGAN PERBAIKAN RAM + GPU
# ===========================

import math
import shutil
import gc

# ============================================================
#   CLEAN MEMORY (RAM + GPU)
# ============================================================
def free_all_memory():
    gc.collect()
    if torch.cuda.is_available():
        torch.cuda.empty_cache()
        torch.cuda.ipc_collect()


# ============================================================
#   RUN SINGLE CONFIG (MEMORY-SAFE VERSION)
# ============================================================
def run_config(run_id: int, cfg: Dict[str, Any], output_root: str, device: torch.device):

    print(f"\n>>> Starting run {run_id} (MEMORY-SAFE)")

    # ===== CLEAN MEMORY BEFORE RUN =====
    free_all_memory()

    # ----- DIRECTORIES -----
    run_dir = os.path.join(output_root, f"run_{run_id}")
    os.makedirs(run_dir, exist_ok=True)
    models_dir = os.path.join(run_dir, "models"); os.makedirs(models_dir, exist_ok=True)
    preds_dir  = os.path.join(run_dir, "preds");  os.makedirs(preds_dir, exist_ok=True)
    train_sample_dir = os.path.join(run_dir, "train_samples"); os.makedirs(train_sample_dir, exist_ok=True)

    # ----- DATASET -----
    ds = SegDataset(cfg["img_dir"], cfg["mask_dir"],
                    img_size=tuple(cfg["img_size"]),
                    augment=cfg.get("augment", False))

    # split
    if cfg.get("use_kfold", False):
        train_idx = cfg["fold_train_idx"]
        test_idx  = cfg["fold_test_idx"]
    else:
        idxs = list(range(len(ds)))
        random.shuffle(idxs)
        ntrain = int(len(ds) * cfg["split_ratio"])
        train_idx = idxs[:ntrain]
        test_idx  = idxs[ntrain:]

    train_set = Subset(ds, train_idx)
    test_set  = Subset(ds, test_idx)

    train_loader = DataLoader(train_set, batch_size=cfg["batch_size"],
                              shuffle=True, num_workers=0)
    test_loader  = DataLoader(test_set, batch_size=1,
                              shuffle=False, num_workers=0)

    # ----- MODEL -----
    model = UNetSmall(
        in_ch=3,
        base_filters=[64, 128],
        bottleneck_filters=cfg["bottleneck_filters"]
    ).to(device).float()

    # ===== Multi-GPU (2 GPU) =====
    if torch.cuda.device_count() >= 2:
        #print("🔥 Using 2 GPUs (DataParallel)")
        model = nn.DataParallel(model, device_ids=[0,1])

    # ----- OPTIMIZER -----
    opt_name = cfg["optimizer"].lower()

    if opt_name == "adamw":
        optimizer = torch.optim.AdamW(
            model.parameters(),
            lr=cfg["learning_rate"],
            weight_decay=cfg.get("weight_decay", 1e-4)
        )

    elif opt_name == "adam":
        optimizer = torch.optim.Adam(
            model.parameters(),
            lr=cfg["learning_rate"],
            weight_decay=cfg.get("weight_decay", 1e-4)
        )

    elif opt_name == "sgd":
        optimizer = torch.optim.SGD(
            model.parameters(),
            lr=cfg["learning_rate"],
            momentum=0.9,
            nesterov=True,
            weight_decay=cfg.get("weight_decay", 1e-4)
        )

    elif opt_name == "rmsprop":
        optimizer = torch.optim.RMSprop(
            model.parameters(),
            lr=cfg["learning_rate"],
            momentum=0.9,
            weight_decay=cfg.get("weight_decay", 1e-4)
        )

    elif opt_name == "radam":
        optimizer = torch.optim.RAdam(
            model.parameters(),
            lr=cfg["learning_rate"],
            weight_decay=cfg.get("weight_decay", 1e-4)
        )

    elif opt_name == "adamax":
        optimizer = torch.optim.Adamax(
            model.parameters(),
            lr=cfg["learning_rate"],
            weight_decay=cfg.get("weight_decay", 1e-4)
        )

    elif opt_name == "adagrad":
        optimizer = torch.optim.Adagrad(
            model.parameters(),
            lr=cfg["learning_rate"],
            weight_decay=cfg.get("weight_decay", 1e-4)
        )

    elif opt_name == "adadelta":
        optimizer = torch.optim.Adadelta(
            model.parameters(),
            lr=cfg["learning_rate"]
        )

    elif opt_name == "lion":
        optimizer = Lion(
            model.parameters(),
            lr=cfg["learning_rate"],
            betas=(0.9, 0.99),
            weight_decay=cfg.get("weight_decay", 1e-4)
        )

    else:
        raise ValueError(f"Unknown optimizer: {cfg['optimizer']}")



    # ----- PSO -----
    real_model = model.module if isinstance(model, nn.DataParallel) else model
    bvec = get_bvec_cpu(real_model)

    pso_obj = None
    if cfg.get("use_pso", True) and bvec.numel() > 0:
        pso_obj = PSO_CPU(
            dim=bvec.numel(),
            n_particles=cfg["pso_particles"],
            w=cfg["pso_w"],
            c1=cfg["pso_c1"],
            c2=cfg["pso_c2"],
            init_pos=bvec
        )

    # ============================================================
    #   TRAINING
    # ============================================================
    train_rows = []

    for epoch in range(1, cfg["epochs"] + 1):

        free_all_memory()  # prevent GPU growth

        t0 = time.time()

        mean_loss, mean_iou, mean_dice, mean_pa, steps = train_one_epoch(
            model=model,
            loader=train_loader,
            optimizer=optimizer,
            device=device,
            tv_weight=cfg.get("tv_weight", 0.0),
            pso_obj=pso_obj,
            pso_trigger_every_batches=cfg.get("pso_trigger_every_batches", 50),
            pso_iterations=cfg.get("pso_iterations", 1),
            cfg=cfg
        )

        t1 = time.time()

        train_rows.append([
            epoch, mean_loss, mean_iou, mean_dice,
            mean_pa, round(t1 - t0, 2)
        ])

        print(f"[RUN {run_id}] Epoch {epoch}/{cfg['epochs']} "
              f"loss={mean_loss:.4f} IoU={mean_iou:.4f} "
              f"Dice={mean_dice:.4f} PA={mean_pa:.4f} "
              f"time={t1 - t0:.2f}s")

        # SAVE SAMPLE
        try:
            sample_imgs, _ = next(iter(train_loader))
            with torch.no_grad():
                pred = torch.sigmoid(model(sample_imgs.to(device)))[0,0].cpu().numpy()
            out_path = os.path.join(train_sample_dir, f"epoch_{epoch}_sample.png")
            save_pred_overlay(sample_imgs[0].cpu(), pred, out_path)
        except:
            pass

    # ============================================================
    #   TESTING
    # ============================================================
    free_all_memory()

    test_rows, test_sum = evaluate_on_loader(
        model=model,
        loader=test_loader,
        device=device,
        postproc=cfg.get("postproc", True),
        binary_threshold=cfg.get("binary_threshold", 0.5),
        min_comp_area=cfg.get("min_component_area", 300),
        save_preds_dir=preds_dir
    )

    # ===========================================
    # HANDLE RUN BURUK — tetap buat excel
    # ===========================================
    if test_sum.get("bad_run", False):

        print(f"⛔ RUN {run_id} DI-SKIP (Prediksi all black/white). "
            "Nilai test = 0 dan tetap disimpan ke Excel.")

        # Bangun summary dummy bernilai 0
        summary = {
            "run_id": run_id,
            "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),

            "img_size": str(cfg["img_size"]),
            "split_ratio": cfg.get("split_ratio", 0.8),
            "num_train_images": 0,
            "num_test_images": 0,

            "optimizer": cfg["optimizer"],
            "learning_rate": cfg["learning_rate"],
            "epochs": cfg["epochs"],
            "batch_size": cfg["batch_size"],
            "loss_fn": "BCE+Dice",
            "bottleneck_filters": cfg["bottleneck_filters"],

            "use_pso": cfg.get("use_pso", False),
            "pso_particles": cfg.get("pso_particles", 0),
            "pso_iterations": cfg.get("pso_iterations", 0),

            "train_mean_loss": 0.0,
            "train_mean_iou": 0.0,
            "train_mean_dice": 0.0,
            "train_mean_pa": 0.0,
            "train_time_total": 0.0,
            "train_time_per_epoch": 0.0,

            "test_mean_iou": 0.0,
            "test_mean_dice": 0.0,
            "test_mean_pa":   0.0,

            "model_path": "SKIPPED",
            "preds_dir": "SKIPPED",
            "train_samples_dir": "SKIPPED",
        }

        # Simpan Excel
        excel_path = os.path.join(output_root, f"run_{run_id}", f"run_{run_id}_results.xlsx")
        save_excel_journal(
            out_path=excel_path,
            analysis_rows=[summary],
            params_by_run={run_id: cfg},
            train_rows_by_run={run_id: []},
            test_rows_by_run={run_id: []}
        )

        # Bersihkan memori
        free_all_memory()

        # KEMBALIKAN tanpa melanjutkan
        return summary, [], [], cfg


    # ============================================================
    #   SAVE MODEL
    # ============================================================
    model_path = os.path.join(models_dir, f"model_run_{run_id}.pth")
    torch.save(model.state_dict(), model_path)

    # ============================================================
    #   BUILD SUMMARY FOR EXCEL
    # ============================================================
    train_mean_loss = float(np.mean([r[1] for r in train_rows]))
    train_mean_iou  = float(np.mean([r[2] for r in train_rows]))
    train_mean_dice = float(np.mean([r[3] for r in train_rows]))
    train_mean_pa   = float(np.mean([r[4] for r in train_rows]))
    train_total_time = float(sum([r[5] for r in train_rows]))

    summary = {
        "run_id": run_id,
        "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),

        "img_size": str(cfg["img_size"]),
        "split_ratio": cfg.get("split_ratio", 0.8),
        "num_train_images": len(train_rows),
        "num_test_images": len(test_rows),

        "optimizer": cfg["optimizer"],
        "learning_rate": cfg["learning_rate"],
        "epochs": cfg["epochs"],
        "batch_size": cfg["batch_size"],
        "loss_fn": cfg.get("loss_fn", "BCE+Dice"),
        "bottleneck_filters": cfg["bottleneck_filters"],

        "use_pso": cfg.get("use_pso", True),
        "pso_particles": cfg.get("pso_particles", 0),
        "pso_iterations": cfg.get("pso_iterations", 0),

        "train_mean_loss": train_mean_loss,
        "train_mean_iou": train_mean_iou,
        "train_mean_dice": train_mean_dice,
        "train_mean_pa": train_mean_pa,
        "train_time_total": train_total_time,
        "train_time_per_epoch": train_total_time / len(train_rows),

        "test_mean_iou": test_sum["mean_iou"],
        "test_mean_dice": test_sum["mean_dice"],
        "test_mean_pa":   test_sum["mean_pa"],

        "model_path": model_path,
        "preds_dir": preds_dir,
        "train_samples_dir": train_sample_dir,
    }

    # ============================================================
    #   SAVE EXCEL PER RUN
    # ============================================================
    excel_path = os.path.join(run_dir, f"run_{run_id}_results.xlsx")

    save_excel_journal(
        out_path=excel_path,
        analysis_rows=[summary],
        params_by_run={run_id: cfg},
        train_rows_by_run={run_id: train_rows},
        test_rows_by_run={run_id: test_rows}
    )

    print(f"📘 Excel saved for RUN {run_id}: {excel_path}")

    # ============================================================
    #   CLEAN EVERYTHING BEFORE RETURN
    # ============================================================
    del model
    del train_loader
    del test_loader
    del train_set
    del test_set
    del ds
    free_all_memory()

    # ============================================================
    #   RETURN 4 VALUES (FIXED)
    # ============================================================
    return summary, train_rows, test_rows, cfg


# -----------------------------
# FINAL: run grid of configs and save excel (MEMORY SAFE VERSION)
# -----------------------------



# =======================
# BLOK 3 / MAIN PROGRAM
# =======================
import datetime
import shutil

if __name__ == "__main__":

    # =====================================================
    # 1. PATH DATASET & OUTPUT
    # =====================================================
    BASE = "/home/jovyan/work/ricardusap/2026_FLOOD/dataset/Flood Area Segmentation"
    IMG_DIR = os.path.join(BASE, "Image")
    MASK_DIR = os.path.join(BASE, "Mask")

    BASEOUTPUT = "/home/jovyan/work/ricardusap/2026_FLOOD/code/unet/PSO_onBottleneck/output"

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    OUTPUT_ROOT = os.path.join(BASEOUTPUT, f"journal_results_{timestamp}")
    os.makedirs(OUTPUT_ROOT, exist_ok=True)

    device = get_device()

    # =====================================================
    # 2. GLOBAL SWITCH PSO
    # =====================================================
    USE_PSO = False      # <<<<< GANTI DI SINI (True/False)

    if USE_PSO:
        print("🔵 PSO MODE ENABLED")
    else:
        print("🟢 PSO MODE DISABLED")


    # =====================================================
    # 3. PARAMETER LISTS
    # =====================================================   
    # ---------- Define parameter lists for analysis ----------
    #optimizer_list = ["AdamW"]
    #optimizer_list = ["Lion","AdamW","Lion","SGD","RMSprop","RAdam","Adamax","Adagrad","Adadelta",]
    optimizer_list = ["AdamW", "Lion", "Adam"]


    learning_rate_list = [1e-3, 1e-4, 5e-4]
    epochs_list = [5, 10, 15, 20, 25, 30, 40, 50, 60, 70, 80, 90, 100]       
    batch_size_list = [64]

    bottleneck_filters_list = [128, 256]
    depth_list = [3]   # currently depth used as metadata; UNetSmall structure fixed

    # ======================================
    # PSO LIST ADAPTIVE (only used if PSO on)
    # ======================================
    if USE_PSO:
        pso_particles_list = [2, 4, 6, 8, 10, 20, 30, 40, 50, 60, 70, 80, 90, 100]
        pso_iterations_list = [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 20, 30, 40, 50, 60, 70, 80, 90, 100]
    else:
        pso_particles_list = [0]      # << prevents huge loops
        pso_iterations_list = [0]

    # =====================================================
    # 4. BUILD GRID SEARCH COMBINATIONS
    # =====================================================
    from itertools import product

    if USE_PSO:
        combos = list(product(
            optimizer_list,
            learning_rate_list,
            epochs_list,
            batch_size_list,
            pso_particles_list,
            pso_iterations_list,
            bottleneck_filters_list,
            depth_list
        ))
    else:
        combos = list(product(
            optimizer_list,
            learning_rate_list,
            epochs_list,
            batch_size_list,
            bottleneck_filters_list,
            depth_list
        ))

    print(f"Total runs to execute: {len(combos)}")

    # =====================================================
    # 5. TRAINING LOOP (ONE RUN PER COMBO)
    # =====================================================
    run_id = 0
    for combo in combos:
        
        # unpack dynamic combo
        if USE_PSO:
            opt, lr, ep, bs, pso_p, pso_it, bf, d = combo
        else:
            opt, lr, ep, bs, bf, d = combo
            pso_p = 0
            pso_it = 0

        run_id += 1
        print(f"\n============================================")
        print(f"RUN {run_id}/{len(combos)}")
        print(f"============================================")

        # =====================================================
        # CONFIG DICTIONARY (PER RUN)
        # =====================================================
        cfg = {
            "img_dir": IMG_DIR,
            "mask_dir": MASK_DIR,
            "img_size": (256, 256),
            "augment": False,

            "optimizer": opt,
            "learning_rate": float(lr),
            "epochs": int(ep),
            "batch_size": int(bs),
            "loss_fn": "BCE+Dice",

            # --- PSO AUTO ---
            "use_pso": USE_PSO,
            "pso_particles": int(pso_p),
            "pso_iterations": int(pso_it),
            "pso_w": 0.5,
            "pso_c1": 1.0,
            "pso_c2": 1.0,
            "pso_trigger_every_batches": 50,
            "use_pso_interval": USE_PSO,

            # --- MODEL ---
            "bottleneck_filters": int(bf),
            "unet_depth": int(d),

            # --- POST PROCESS ---
            "postproc": True,
            "binary_threshold": 0.5,
            "min_component_area": 100,
            "tv_weight": 0.0,

            # --- SPLIT ---
            "split_ratio": 0.8,
            "use_kfold": False
        }

        # RUN TRAIN + TEST
        summary, train_rows, test_rows, full_cfg = run_config(
            run_id, cfg, OUTPUT_ROOT, device
        )

        # =====================================================
        # 6. SAVE EXCEL PER RUN (MEMORY SAFE)
        # =====================================================
        excel_out = os.path.join(
            OUTPUT_ROOT, f"run_{run_id}", 
            f"run_{run_id}_results.xlsx"
        )

        save_excel_journal(
            excel_out,
            analysis_rows=[summary],
            params_by_run={run_id: full_cfg},
            train_rows_by_run={run_id: train_rows},
            test_rows_by_run={run_id: test_rows},
        )

        print(f"Excel saved to: {excel_out}")

    # =====================================================
    # END
    # =====================================================
    #print("\nAll experiments complete.")
    print(f"Results saved in: {OUTPUT_ROOT}")


